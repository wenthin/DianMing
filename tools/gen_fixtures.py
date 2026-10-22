#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成单元测试用的表格样本：xlsx / xls / csv(GBK)"""
import os
import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "app", "src", "test", "resources")
os.makedirs(OUT, exist_ok=True)

HEADERS = ["学号", "姓名", "性别", "手机", "入学日期", "语文", "英语"]
ROWS = [
    ["20230101", "张三", "男", "13800138000", datetime.datetime(2023, 9, 1), 88, 92],
    ["20230102", "李四", "女", "13900139000", datetime.datetime(2023, 9, 1), 95, 97],
    ["20230103", "王五", "男", "13611136111", datetime.datetime(2023, 9, 2), 76, 81],
    ["20230104", "赵六", "女", "13722237222", datetime.datetime(2023, 9, 2), 90, 85],
    ["20230105", "孙七", "男", "", datetime.datetime(2023, 9, 3), 66, 73],
]

# ---------------- xlsx ----------------
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "学生名单"
ws.append(HEADERS)
for r in ROWS:
    ws.append(r)
wb.save(os.path.join(OUT, "roster.xlsx"))

# 再来一个"第一行不是表头"的 xlsx（表头在第二行）
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "Sheet1"
ws2.append(["XX中学高一新生名单"])
ws2.append(HEADERS)
for r in ROWS:
    ws2.append(r)
wb2.save(os.path.join(OUT, "roster_with_title.xlsx"))

# ---------------- xls ----------------
import xlwt
book = xlwt.Workbook(encoding="utf-8")
sheet = book.add_sheet("学生名单")
date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
for c, h in enumerate(HEADERS):
    sheet.write(0, c, h)
for r, row in enumerate(ROWS, start=1):
    for c, v in enumerate(row):
        if isinstance(v, datetime.datetime):
            sheet.write(r, c, v, date_style)
        else:
            sheet.write(r, c, v)
book.save(os.path.join(OUT, "roster.xls"))

# ---------------- csv (GBK 编码，模拟国内 Excel 导出) ----------------
with open(os.path.join(OUT, "roster.csv"), "wb") as f:
    f.write((",".join(HEADERS) + "\n").encode("gbk"))
    for row in ROWS:
        vals = [x.strftime("%Y-%m-%d") if isinstance(x, datetime.datetime) else str(x) for x in row]
        f.write((",".join(vals) + "\n").encode("gbk"))

print("generated fixtures in", OUT)
for n in sorted(os.listdir(OUT)):
    print(" -", n, os.path.getsize(os.path.join(OUT, n)), "bytes")
